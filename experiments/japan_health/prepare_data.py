from auditableai.paths import REPO, WORK, JAPAN, EVIDENCE, FIGURES, TABLES, CHECKS, display_path, resolve_record_path, prepare_outputs, required_input
"""Read-only extraction of official tables with cell-level source coordinates."""
from pathlib import Path
import json, hashlib, re, sys
import openpyxl
import pandas as pd

WS = WORK
E = JAPAN
R = E/'raw'
P = E/'processed'
AGES = list(range(40,75,5))
CATEGORIES = ['90.0以上','85.0以上90.0未満','85.0未満']

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()

def extract():
    P.mkdir(parents=True,exist_ok=False)
    allrows=[]; unknown=[]; checks=[]; prefectures=None
    for release in range(3,12):
        path=R/f'ndb{release}_waist.xlsx'
        sheet=openpyxl.load_workbook(path,read_only=True,data_only=True).active
        rows=list(sheet.values)
        assert rows[2][2]=='男' and rows[2][10]=='女'
        assert list(rows[3][2:9])==[f'{a}～{a+4}歳' for a in AGES]
        assert list(rows[3][10:17])==list(rows[3][2:9])
        year=2012+release
        assert (str(year) in rows[0][0]) or (f'H{year-1988}' in rows[0][0])
        groups={}; pref=None
        for ri,row in enumerate(rows[5:],6):
            if row[0]: pref=row[0]
            if row[1]:groups.setdefault(pref,[]).append((ri,row))
        unknown_group=groups.pop('都道府県判別不可',[])
        assert len(groups)==47
        if prefectures is None:prefectures=list(groups)
        assert list(groups)==prefectures
        subtotal_checks=0
        for pi,(pref,group) in enumerate(groups.items(),1):
            assert [x[1][1] for x in group]==CATEGORIES
            for sex,start in [('male',2),('female',10)]:
                for ri,row in group:
                    if all(isinstance(v,(int,float)) for v in row[start:start+8]):
                        assert sum(row[start:start+7])==row[start+7],(year,pref,sex,ri)
                        subtotal_checks+=1
                for ai,age in enumerate(AGES):
                    counts=[x[1][start+ai] for x in group]
                    complete=all(isinstance(v,(int,float)) and v>=0 for v in counts)
                    n=sum(counts) if complete else None
                    y=sum(counts[:2]) if complete and sex=='male' else counts[0] if complete else None
                    coords=[f'{openpyxl.utils.get_column_letter(start+ai+1)}{x[0]}' for x in group]
                    allrows.append(dict(year=year,release=release,pref_id=pi,prefecture=pref,sex=sex,age=age,
                        n=n,y=y,rate=y/n if complete and n>0 else None,complete=complete,
                        source_file=path.name,sheet=sheet.title,source_cells=';'.join(coords),source_sha256=sha(path)))
        for ri,row in unknown_group:unknown.append(dict(year=year,source_file=path.name,row=ri,values=list(row)))
        checks.append(dict(year=year,release=release,title=rows[0][0],prefectures=47,
                           subtotal_equalities=subtotal_checks,unknown_rows=len(unknown_group),sha256=sha(path)))
    df=pd.DataFrame(allrows)
    assert not df.duplicated(['year','pref_id','sex','age']).any()
    df.to_csv(P/'waist_panel.csv',index=False)
    (P/'unknown_prefecture_rows.json').write_text(json.dumps(unknown,ensure_ascii=False,indent=2))
    for c in checks:
        d=df[df.year==c['year']];c['complete_strata']=int(d.complete.sum());c['strata']=len(d)
    (P/'extraction_checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
    path=R/'census2020_table2_7.xlsx'
    sheet=openpyxl.load_workbook(path,read_only=True,data_only=True).active
    weights=[]
    for ri,row in enumerate(sheet.values,1):
        if row[0]=='0_国籍総数' and row[7]=='00000' and row[1] in ['1_男','2_女']:
            sex='male' if row[1]=='1_男' else 'female'
            counts=list(row[18:25]);assert all(isinstance(v,(int,float)) for v in counts)
            for ai,(age,n) in enumerate(zip(AGES,counts)):
                weights.append(dict(sex=sex,age=age,population=n,weight=n/sum(counts),
                    source_file=path.name,sheet=sheet.title,source_cell=f'{openpyxl.utils.get_column_letter(19+ai)}{ri}',source_sha256=sha(path)))
    assert len(weights)==14,weights
    pd.DataFrame(weights).to_csv(P/'census_weights.csv',index=False)
    hist=df[df.year==2019].groupby(['sex','age']).n.sum().reset_index()
    hist['weight']=hist.n/hist.groupby('sex').n.transform('sum')
    hist.to_csv(P/'ndb2019_weights.csv',index=False)
    print(json.dumps(checks,ensure_ascii=False,indent=2))
    print(pd.DataFrame(weights)[['sex','age','population','weight']].to_string(index=False))

if __name__=='__main__':extract()
