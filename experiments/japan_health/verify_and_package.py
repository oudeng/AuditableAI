from auditableai.paths import REPO, WORK, JAPAN, EVIDENCE, FIGURES, TABLES, CHECKS, display_path, resolve_record_path, prepare_outputs, required_input
"""Independent raw-cell calculations, retained-model replay, and claim records.

Normal reproduction tasks compare two fully equipped programmatic routes. They
are not a user study, a fault-detection benchmark, or an audit-superiority test.
"""
from pathlib import Path
import json, time, pickle, hashlib, sys
import numpy as np
import pandas as pd
import openpyxl
from scipy.special import expit,logit
from run_formal import WS,E,P,HERE,FEATURES,FAMILIES,make_panel,predict,aggregate,metrics,sha,save,native_predict

O=E/'formal_v1'

def source_oracle():
    # A separate extraction path uses direct Excel indexing and finds headers.
    census=openpyxl.load_workbook(E/'raw/census2020_table2_7.xlsx',read_only=True,data_only=True).active
    rows=list(census.values);head=rows[9]
    weights={}
    for row in rows[12:]:
        if row[0]=='0_国籍総数' and row[7]=='00000' and row[1] in ['1_男','2_女']:
            sex='male' if row[1]=='1_男' else 'female'
            values=[row[next(i for i,x in enumerate(head) if isinstance(x,str) and x.endswith(f'{a}～{a+4}歳'))] for a in range(40,75,5)]
            weights[sex]=np.array(values)/sum(values)
    rates={};counts={};cells={};metadata={}
    for year in range(2015,2024):
        sheet=openpyxl.load_workbook(E/f'raw/ndb{year-2012}_waist.xlsx',read_only=False,data_only=True).active
        names=[(i,sheet.cell(i,1).value) for i in range(6,sheet.max_row+1) if sheet.cell(i,1).value not in [None,'都道府県判別不可']]
        assert len(names)==47
        for pref,(start,name) in enumerate(names,1):
            for sex,first in [('male',3),('female',11)]:
                ny=[];ns=[];coords=[]
                for age,col in zip(range(40,75,5),range(first,first+7)):
                    vals=[sheet.cell(start+k,col).value for k in range(3)]
                    assert all(isinstance(v,(int,float)) for v in vals)
                    ns.append(sum(vals));ny.append(sum(vals[:2]) if sex=='male' else vals[0])
                    coords.append([sheet.cell(start+k,col).coordinate for k in range(3)])
                key=(year,pref,sex)
                metadata[key]={'year':year,'release':year-2012,'sex':sex,'threshold_cm':85 if sex=='male' else 90,'source_file':f'ndb{year-2012}_waist.xlsx','sheet':sheet.title}
                assert str(year) in str(sheet.cell(1,1).value) or f'H{year-1988}' in str(sheet.cell(1,1).value)
                counts[key]=(ny,ns);rates[key]=float(np.dot(weights[sex],np.array(ny)/ns));cells[key]=coords
    return rates,counts,cells,metadata

def main():
    freeze=json.loads((O/'freeze.json').read_text())
    assert all(sha(resolve_record_path(p))==v for p,v in freeze['files'].items())
    oracle,counts,cells,metadata=source_oracle()
    df=pd.read_csv(P/'waist_panel.csv')
    for source,d in df.groupby('source_file'):
        assert d.source_sha256.nunique()==1 and sha(E/'raw'/source)==d.source_sha256.iloc[0]
    panel=make_panel();test=panel[panel.year>=2022]
    selection=json.loads((O/'selection.json').read_text());spatial=json.loads((O/'spatial_selection.json').read_text())
    predictions=pd.read_csv(O/'predictions.csv')
    model_replay={}
    for fam in FAMILIES:
        m=pickle.loads((O/f'{fam}.pkl').read_bytes())
        p=predict(fam,selection[fam]['params'],m,test)
        a=aggregate(test,p)
        retained=predictions[(predictions.model==fam)&(predictions.evaluation=='temporal')]
        joined=a.merge(retained,on=['year','pref_id','prefecture','sex'],suffixes=('_new','_saved'))
        err=float(np.max(np.abs(joined.predicted_new-joined.predicted_saved)))
        assert err<1e-12;model_replay[fam]=err
        trace=selection[fam]['selection'];assert trace['selected_index']==int(np.argmin([c['mean_mae_pp'] for c in trace['candidates']]))
    protocol=json.loads((HERE/'protocol.json').read_text())
    for region,ids in protocol['spatial_groups'].items():
        for fam in FAMILIES:
            m=pickle.loads((O/f'spatial_{region}_{fam}.pkl').read_bytes());d=test[test.pref_id.isin(ids)]
            par=spatial[region][fam]['params'];a=aggregate(d,predict(fam,par,m,d))
            b=predictions[(predictions.model==fam)&(predictions.evaluation=='spatial')&predictions.pref_id.isin(ids)]
            j=a.merge(b,on=['year','pref_id','sex'],suffixes=('_new','_saved'))
            err=float(np.max(np.abs(j.predicted_new-j.predicted_saved)));assert err<1e-12
            model_replay[f'{region}_{fam}']=err
            tr=spatial[region][fam]['selection'];assert tr['selected_index']==int(np.argmin([c['mean_mae_pp'] for c in tr['candidates']]))
    raw_error=0.
    for r in df.itertuples():
        y,n=counts[(r.year,r.pref_id,r.sex)];k=(r.age-40)//5
        assert (r.y,r.n)==(y[k],n[k])
    for r in predictions.itertuples():raw_error=max(raw_error,abs(r.observed-oracle[(r.year,r.pref_id,r.sex)]))
    assert raw_error<1e-12
    claims=[]
    for (year,pref,sex),d in test.groupby(['year','pref_id','sex'],sort=True):
        d=d.sort_values('age');key=(year,pref,sex)
        outputs={fam:float(predictions[(predictions.evaluation=='temporal')&(predictions.model==fam)&(predictions.year==year)&(predictions.pref_id==pref)&(predictions.sex==sex)].iloc[0].predicted) for fam in FAMILIES}
        claims.append({'claim_id':f'JP-{year}-{pref:02d}-{sex}',
            'assertion':{'indicator':'waist criterion proportion','observed_standardized':oracle[key],'model_forecasts':outputs},
            'scope':{'year':int(year),'release':int(year-2012),'pref_id':int(pref),'prefecture':d.iloc[0].prefecture,'sex':sex,'ages':'40–74 fiscal-year-end','unit':'proportion','threshold_cm':85 if sex=='male' else 90,'population':'waist examinees, not all residents'},
            'artifact':{'model_hashes':{fam:sha(O/f'{fam}.pkl') for fam in FAMILIES},'gate_expression_sha256':sha(O/'gate_expression.json')},
            'reference':{'source_file':d.iloc[0].source_file,'sheet':d.iloc[0].sheet,'sha256':d.iloc[0].source_sha256,'category_cells':d.source_cells.tolist(),'weights_file':'processed/census_weights.csv','weights_sha256':sha(P/'census_weights.csv')},
            'protocol':{'id':protocol['id'],'sha256':sha(HERE/'protocol.json'),'information':'lag1 and lag2 only; final fit through FY2021'},
            'evidence':{'age_numerators':d.y.astype(int).tolist(),'age_denominators':d.n.astype(int).tolist(),'weights':d.weight.tolist(),'source_cell_groups':cells[key]},
            'boundary':'Aggregate examinee surveillance; no individual diagnosis, treatment effect or demonstrated human audit efficiency.'})
    save(O/'claim_records.json',claims)
    # Independent metadata and ordinary files are available to both routes.
    b1={};full={}
    for c in claims:
        s=c['scope'];key=(s['year'],s['pref_id'],s['sex'])
        t=time.perf_counter()
        d=test[(test.year==key[0])&(test.pref_id==key[1])&(test.sex==key[2])].sort_values('age')
        b1[key]={'y':d.y.astype(int).tolist(),'n':d.n.astype(int).tolist(),'rate':float(np.dot(d.weight,d.rate)),'cells':d.source_cells.tolist()}
        b1[key]['metadata']={'year':int(d.iloc[0].year),'release':int(d.iloc[0].release),'sex':key[2],'threshold_cm':85 if key[2]=='male' else 90,'source_file':d.iloc[0].source_file,'sheet':d.iloc[0].sheet}
        b1[key]['seconds']=time.perf_counter()-t
        t=time.perf_counter();ev=c['evidence']
        full[key]={'y':ev['age_numerators'],'n':ev['age_denominators'],'rate':float(np.dot(ev['weights'],np.array(ev['age_numerators'])/ev['age_denominators'])),'cells':c['reference']['category_cells']}
        full[key]['metadata']={**{k:s[k] for k in ['year','release','sex','threshold_cm']},**{k:c['reference'][k] for k in ['source_file','sheet']}}
        full[key]['seconds']=time.perf_counter()-t
    export=json.loads((O/'gate_expression.json').read_text());gate=pickle.loads((O/'gate.pkl').read_bytes())
    err=float(np.max(np.abs(expit(logit(test.lag1.to_numpy())+native_predict(export,test[FEATURES].to_numpy()))-predict('gate',selection['gate']['params'],gate,test))))
    altw=pd.read_csv(P/'ndb2019_weights.csv')
    for sex in ['male','female']:
        ref=np.sum([counts[(2019,pref,sex)][1] for pref in range(1,48)],axis=0).astype(float);ref/=ref.sum()
        assert np.max(np.abs(ref-altw[altw.sex==sex].sort_values('age').weight.to_numpy()))<1e-14
    tasks=[]
    for route,answers in [('B1',b1),('Full',full)]:
        t1=all(v['y']==counts[k][0] and v['n']==counts[k][1] for k,v in answers.items())
        t3=max(abs(v['rate']-oracle[k]) for k,v in answers.items())
        t5=all(v['cells']==[';'.join(x) for x in cells[k]] for k,v in answers.items())
        # Recalculate alternate reference by both connected records and flat files.
        maxalt=0.
        for c in claims:
            s=c['scope'];key=(s['year'],s['pref_id'],s['sex']);w=altw[altw.sex==key[2]].sort_values('age').weight.to_numpy()
            v=answers[key];value=float(np.dot(w,np.array(v['y'])/v['n']))
            y,n=counts[key];ref=float(np.dot(w,np.array(y)/n));maxalt=max(maxalt,abs(value-ref))
        # Release/year and age metadata matched against official headers by extractor;
        # both routes carry the same declared scope, never inferred from model output.
        t2=all(v['metadata']==metadata[k] for k,v in answers.items())
        for task,ok,discrepancy in [('T1',t1,0.),('T2',t2,0.),('T3',t3<1e-10,t3),('T4',err<1e-10,err),('T5',t5,0.),('T6',maxalt<1e-10,maxalt)]:
            tasks.append(dict(route=route,task=task,passed=bool(ok),max_abs_discrepancy=discrepancy))
    assert all(t['passed'] for t in tasks)
    save(O/'reproduction_tasks.json',{'tasks':tasks,'interpretation':'Both scripted routes have complete access and pass. No incremental success advantage, human completion-time evidence, or real-world failure rate is claimed. Source parsing and package creation are excluded from the lookup timing.',
        'warm_lookup_seconds_188_records':{'B1':sum(v['seconds'] for v in b1.values()),'Full':sum(v['seconds'] for v in full.values())},
        'timing_design':'One deterministic lookup batch on this host, not randomized or repeated; diagnostic only, not a comparative performance claim.'})
    save(O/'verification.json',dict(raw_category_cells_read=len(df)*3,normalized_numerator_denominator_pairs_verified=len(df),raw_rate_max_abs_error=raw_error,selected_model_replay=model_replay,
        gate_native_replay_max_abs_error=err,claims=len(claims),tasks_passed=12,
        limitations=['same-environment replay, not independent investigator replication','source-oracle implementation by same assistant','no user study','public revised vintages rather than full vintage reconstruction']))
    save(O/'integrity.json',{p.name:sha(p) for p in sorted(O.iterdir()) if p.is_file() and p.name!='integrity.json'})
    print('Verified',len(df),'strata;',len(model_replay),'retained models;',len(claims),'claims; all 12 task-route checks pass.')

if __name__=='__main__':main()
